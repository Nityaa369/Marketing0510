"""Write data/icps_template.xlsx: 20 EXAMPLE NCA rows plus a READ_ME sheet.

The readers are NCA segments from the project's ICP work; every line of copy is
EXAMPLE copy for testing the pipeline and must be approved before any launch.
Run:  python scripts/make_sample_sheet.py
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from studio.config import ROOT  # noqa: E402
from studio.sheet import COLUMNS  # noqa: E402

CAMPAIGN = "EXAMPLE_NCA_31Oct-2Nov"
S1 = "Qualify through the NCA from home. Become eligible for the Canadian bar."
S2 = "The NCA route takes you there, from India."
S3 = "Follow the NCA route, give exams online, before you land."

# id, icp, headline, supporting, hero scene
ROWS = [
    ("nca01_new_grads", "New law grads",
     "You can qualify NCA, then the *Canadian bar*, while your batchmates are still deciding.", S1,
     "a 23 year old Indian woman law graduate in a white shirt and black blazer holding her degree, Toronto skyline softly visible through a window behind her"),
    ("nca02_final_year", "Final year law students",
     "Convocation se pehle hi, *Canada ki practice* karo plan.", S3,
     "a 21 year old Indian male law student in a college library with a maple leaf bookmark in a law textbook, warm afternoon light"),
    ("nca03_canadian_clients", "Lawyers with Canadian clients",
     "You can qualify NCA, then the *Canadian bar*, and advise your Canadian clients yourself.", S1,
     "a confident Indian lawyer in his 30s on a video call in a modern Delhi office, a client in Toronto visible on the laptop screen as a blurred figure"),
    ("nca04_nri_clients", "Lawyers with NRI clients",
     "You can qualify NCA, then the *Canadian bar*, and serve your NRI clients on both sides.", S2,
     "a 40 year old Indian woman advocate in court dress in her chamber in Chandigarh, case files on the desk, a small Canadian flag pin on her blazer"),
    ("nca05_spouse_in_canada", "Lawyers whose spouse is in Canada",
     "Your spouse is in Canada. You can qualify NCA, then the *Canadian bar*, and practise there too.", S3,
     "an Indian woman lawyer in her late 20s at an airport departure gate with a laptop bag, looking hopeful, maple leaf motif on a distant window"),
    ("nca06_family_in_canada", "Lawyers with family in Canada",
     "Half your family is in Canada. You can qualify NCA, then the *Canadian bar*, from home.", S1,
     "a warm Indian family video call, an Indian male lawyer in his 30s in a black coat smiling at relatives on a tablet, autumn maple trees visible on the screen"),
    ("nca07_punjab_lawyers", "Punjab lawyers",
     "You can qualify NCA, then the *Canadian bar*, and practise where your clients already live.", S2,
     "a Sikh advocate in his 30s with a navy turban and black coat standing outside a district court in Punjab, golden hour"),
    ("nca08_lpo_lawyers", "LPO lawyers",
     "You already draft for foreign firms. You can qualify NCA, then the *Canadian bar*, and sign your own work.", S1,
     "an Indian woman lawyer in her 20s reviewing contracts on two monitors in a bright LPO office in Gurugram at night"),
    ("nca09_in_house", "In-house counsel",
     "You can qualify NCA, then the *Canadian bar*, and take your in-house career to Canada.", S3,
     "a 35 year old Indian in-house counsel in a grey suit in a glass meeting room in Bengaluru, a Toronto skyline poster on the wall behind"),
    ("nca10_gulf_lawyers", "Indian lawyers in the Gulf",
     "You moved once for your career. You can qualify NCA, then the *Canadian bar*, from Dubai.", S1,
     "an Indian male lawyer in his 30s in a crisp suit on a Dubai office balcony at dusk, holding a coffee, calm and ambitious"),
    ("nca11_corporate_associates", "Corporate associates",
     "You can qualify NCA, then the *Canadian bar*, without pausing your practice.", S3,
     "a young Indian woman corporate lawyer in a Mumbai law firm conference room late evening, deal documents and a laptop, city lights outside"),
    ("nca12_llm_planners", "Lawyers planning a Canadian LLM",
     "Before you spend on an LLM, see the NCA route to the *Canadian bar*.", S2,
     "an Indian male lawyer in his 20s comparing university brochures at a desk at home, a maple leaf mug beside him"),
    ("nca13_ielts_ready", "Lawyers with IELTS General done",
     "Your IELTS is done. You can qualify NCA, then the *Canadian bar*, next.", S1,
     "a happy Indian woman lawyer in her 20s holding a result envelope in a sunlit Pune apartment, black blazer over a kurta"),
    ("nca14_contract_lawyers", "Contract lawyers",
     "You can qualify NCA, then the *Canadian bar*, and turn contract work into Canadian Legal Practice.", S3,
     "an Indian male contract lawyer in his 30s working on a laptop in a quiet co-working space in Hyderabad, legal pad with notes"),
    ("nca15_young_litigators", "Young litigators",
     "You can qualify NCA, then the *Canadian bar*, while you keep arguing in Indian courts.", S2,
     "a 27 year old Indian woman advocate in a black gown and white bands walking up the steps of a high court, determined expression"),
    ("nca16_kids_in_canada", "Lawyers whose children study in Canada",
     "Your child studies in Canada. You can qualify NCA, then the *Canadian bar*, and build your practice near them.", S1,
     "a 50 year old Indian senior advocate in court dress in a wood panelled chamber holding a framed photo of his daughter at a Canadian university"),
    ("nca17_llb_in_canada", "Indian law grads living in Canada",
     "Your Indian LLB came with you to Canada. The NCA route takes it to the *Canadian bar*.", S3,
     "an Indian woman in her late 20s with a law book in a Toronto cafe in autumn, red maple leaves outside the window"),
    ("nca18_llm_in_india", "LLM grads in India",
     "You did the LLM. You can qualify NCA, then the *Canadian bar*, with the degree you already hold.", S2,
     "an Indian male LLM graduate in his late 20s in a convocation gown on a campus lawn, holding a scroll, soft daylight"),
    ("nca19_siblings_in_canada", "Lawyers with siblings in Canada",
     "Your brother made it to Canada. You can qualify NCA, then the *Canadian bar*, and join him as a lawyer.", S1,
     "two Indian siblings hugging at an Indian airport arrivals hall, the younger one a lawyer in a black blazer, joyful reunion"),
    ("nca20_senior_advocates", "Senior advocates with Canadian clients",
     "Your Canadian clients trust you already. You can qualify NCA, then the *Canadian bar*, and keep them.", S3,
     "a 55 year old Indian woman senior advocate in a black gown in a grand library, reading glasses, dignified, a Canadian law report on the table"),
]

HEADERS = list(COLUMNS)


def main(out: Path = ROOT / "data" / "icps_template.xlsx") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "ICPs"
    ws.append(HEADERS + ["example_label"])
    for rid, icp, headline, supporting, scene in ROWS:
        rec = {
            "id": rid, "campaign": CAMPAIGN, "icp": icp, "banner": "", "headline": headline,
            "supporting": supporting, "cta": "JOIN NOW", "price": "", "price_line": "",
            "journey_tab": "", "hero_prompt": scene, "hero_image": "", "seed": "",
        }
        ws.append([rec[h] for h in HEADERS] + ["EXAMPLE COPY: approve before launch"])
    bold = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="C8102E")
    widths = {"A": 26, "B": 24, "C": 34, "D": 20, "E": 70, "F": 60, "G": 12, "K": 90, "N": 34}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"

    info = wb.create_sheet("READ_ME")
    info.append(["EXAMPLE DATA. Readers come from the NCA ICP work; all copy is sample copy to test the pipeline."])
    info.append(["Approve every line against nca-ad-rules and nca-claim-rules before launch."])
    info.append([])
    info.append(["column", "what to put"])
    for name, desc in COLUMNS.items():
        info.append([name, desc])
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 100
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(main())
